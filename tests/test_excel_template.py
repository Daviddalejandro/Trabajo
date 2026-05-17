"""Pruebas del constructor de plantilla Excel oficial."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from bpmn_platform.core.catalogs import load_catalogs
from bpmn_platform.excel.schema import build_default_schema
from bpmn_platform.excel.template_builder import build_template


def test_template_contains_expected_sheets(tmp_path: Path) -> None:
    catalogs = load_catalogs()
    schema = build_default_schema()
    out = tmp_path / "plantilla.xlsx"
    build_template(out, schema=schema, catalogs=catalogs, activity_rows=10)

    wb = load_workbook(out)
    assert set(wb.sheetnames) == {
        schema.instructions_sheet,
        schema.process_sheet,
        schema.activities_sheet,
        schema.catalogs_sheet,
    }

    activities = wb[schema.activities_sheet]
    headers = [cell.value for cell in next(activities.iter_rows(min_row=1, max_row=1))]
    expected = [col.header for col in schema.activity_columns]
    assert headers == expected


def test_template_has_data_validations(tmp_path: Path) -> None:
    out = tmp_path / "plantilla.xlsx"
    build_template(out, activity_rows=5)
    wb = load_workbook(out)
    activities = wb["Actividades"]
    # Debe haber al menos una validacion (dropdown Tipo BPMN, Sistema, Roles, etc.).
    assert activities.data_validations.dataValidation


def test_catalog_sheet_lists_all_catalogs(tmp_path: Path) -> None:
    catalogs = load_catalogs()
    out = tmp_path / "plantilla.xlsx"
    build_template(out, catalogs=catalogs, activity_rows=5)
    wb = load_workbook(out)
    catalogs_sheet = wb["Catalogos"]
    headers = [cell.value for cell in catalogs_sheet[4] if cell.value]
    for name in catalogs.all():
        assert name in headers
