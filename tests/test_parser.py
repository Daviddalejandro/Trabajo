"""Tests del parser de Excel (Fase 4)."""
from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook

from bpmn_platform.core.catalogs import load_catalogs
from bpmn_platform.excel.parser import ExcelParser
from bpmn_platform.excel.schema import build_default_schema
from bpmn_platform.excel.template_builder import build_template


def _fill_template(
    path: Path,
    *,
    process_row: Sequence,
    activity_rows: Sequence[Sequence],
) -> None:
    schema = build_default_schema()
    catalogs = load_catalogs()
    build_template(path, schema=schema, catalogs=catalogs, activity_rows=20)
    wb = load_workbook(path)
    proc = wb[schema.process_sheet]
    for col, value in enumerate(process_row, start=1):
        proc.cell(row=3, column=col, value=value)
    acts = wb[schema.activities_sheet]
    for r, row in enumerate(activity_rows, start=3):
        for col, value in enumerate(row, start=1):
            acts.cell(row=r, column=col, value=value)
    wb.save(path)


def test_parser_happy_path(tmp_path: Path) -> None:
    template = tmp_path / "plantilla.xlsx"
    process_row = (
        "PROC-001",
        "Onboarding clientes",
        "Comercial",
        "USUARIO_NEGOCIO",
        "1.0.0",
        "Alta de cliente nuevo.",
    )
    activity_rows = [
        ("ACT-001", "PROC-001", "StartEvent", "Iniciar onboarding"),
        (
            "ACT-002",
            "PROC-001",
            "UserTask",
            "Validar identidad cliente",
            "ANALISTA_RIESGO",
            "CRM",
            None,
            None,
            "Documento identidad",
            "Identidad validada",
            None,
            "R-SUPLANTACION",
            "C-BIOMETRICO",
            "2h",
            "Tasa validacion",
        ),
        (
            "ACT-003",
            "PROC-001",
            "ServiceTask",
            "Generar reporte regulatorio",
            "SISTEMA",
            "AIRFLOW",
            "python reconcile.py",
            None,
            None,
            None,
            None,
            None,
            None,
            "1h",
        ),
        ("ACT-004", "PROC-001", "EndEvent", "Finalizar onboarding"),
    ]
    _fill_template(template, process_row=process_row, activity_rows=activity_rows)

    result = ExcelParser().parse(template)

    assert result.ok, [f"{i.code}: {i.message}" for i in result.errors]
    model = result.model
    assert len(model.processes) == 1
    assert model.processes[0].id == "PROC-001"
    assert len(model.events) == 2
    assert len(model.activities) == 2
    assert {a.id for a in model.activities} == {"ACT-002", "ACT-003"}
    assert len(model.commands) == 1
    cmd = model.commands[0]
    assert cmd.application_id == "AIRFLOW"
    assert cmd.invocation == "python reconcile.py"
    assert any(use.command_ids for use in model.activity_application_uses)
    assert any(s.id == "SLA-ACT-002" for s in model.slas)
    assert any(asset.name == "Documento identidad" for asset in model.information_assets)
    assert any(k.name == "Tasa validacion" for k in model.kpis)


def test_parser_detects_orphan_command(tmp_path: Path) -> None:
    template = tmp_path / "plantilla.xlsx"
    _fill_template(
        template,
        process_row=("PROC-1", "Proceso X"),
        activity_rows=[
            (
                "ACT-1",
                "PROC-1",
                "ServiceTask",
                "Ejecutar batch X",
                "SISTEMA",
                None,
                "python bad.py",
            ),
        ],
    )
    result = ExcelParser().parse(template)
    codes = {i.code for i in result.issues}
    assert "CMD-ORPHAN" in codes
    assert not result.ok


def test_parser_detects_invalid_bpmn_type(tmp_path: Path) -> None:
    template = tmp_path / "plantilla.xlsx"
    _fill_template(
        template,
        process_row=("PROC-1", "Proceso X"),
        activity_rows=[
            ("ACT-1", "PROC-1", "FooBar", "Validar identidad"),
        ],
    )
    result = ExcelParser().parse(template)
    codes = {i.code for i in result.issues}
    assert "BPMN-TYPE-INVALID" in codes


def test_parser_detects_duplicate_activity_ids(tmp_path: Path) -> None:
    template = tmp_path / "plantilla.xlsx"
    _fill_template(
        template,
        process_row=("PROC-1", "Proceso X"),
        activity_rows=[
            ("ACT-1", "PROC-1", "UserTask", "Validar identidad", "ANALISTA_RIESGO"),
            ("ACT-1", "PROC-1", "UserTask", "Validar identidad", "ANALISTA_RIESGO"),
        ],
    )
    result = ExcelParser().parse(template)
    codes = {i.code for i in result.issues}
    assert "ACT-ID-DUPLICATE" in codes


def test_parser_accepts_business_labels(tmp_path: Path) -> None:
    """Permitir que el usuario diligencie con labels en espanol."""
    template = tmp_path / "plantilla.xlsx"
    _fill_template(
        template,
        process_row=("PROC-1", "Proceso X"),
        activity_rows=[
            ("EVT-START", "PROC-1", "Inicio", "Iniciar proceso"),
            ("ACT-1", "PROC-1", "Actividad Humana", "Validar identidad cliente", "ANALISTA_RIESGO"),
            ("EVT-END", "PROC-1", "Fin", "Finalizar proceso"),
        ],
    )
    result = ExcelParser().parse(template)
    assert result.ok, [f"{i.code}: {i.message}" for i in result.errors]
    assert len(result.model.activities) == 1
    assert len(result.model.events) == 2


def test_parser_warns_unknown_catalog_values(tmp_path: Path) -> None:
    template = tmp_path / "plantilla.xlsx"
    _fill_template(
        template,
        process_row=("PROC-1", "Proceso X"),
        activity_rows=[
            (
                "ACT-1",
                "PROC-1",
                "UserTask",
                "Validar identidad cliente",
                "ROL_INEXISTENTE",
                "SISTEMA_DESCONOCIDO",
            ),
        ],
    )
    result = ExcelParser().parse(template)
    codes = {i.code for i in result.issues}
    assert "ACT-ROLE-UNKNOWN" in codes
    assert "ACT-SYSTEM-UNKNOWN" in codes
