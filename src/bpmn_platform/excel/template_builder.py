"""Constructor de la plantilla Excel oficial (Fase 3).

Genera un archivo .xlsx con:
  * Hoja `Instrucciones` (reglas para usuarios de negocio).
  * Hoja `Proceso`     (metadata del proceso).
  * Hoja `Actividades` (filas BPMN con dropdowns y validaciones).
  * Hoja `Catalogos`   (codigos y etiquetas que alimentan los dropdowns).

Los dropdowns referencian rangos de la hoja `Catalogos`, por lo que el
archivo sigue siendo valido aunque se abra fuera de la app.
"""
from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..config import settings
from ..core.catalogs import CatalogBundle, load_catalogs
from ..core.catalogs.loader import Catalog
from ..logging_config import get_logger
from .schema import CellKind, ColumnSpec, ExcelSchema, build_default_schema

log = get_logger(__name__)


# --- Estilos -------------------------------------------------------------- #

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_REQUIRED_FONT = Font(color="C00000", bold=True, size=11)
_HELP_FONT = Font(color="595959", italic=True, size=9)
_TITLE_FONT = Font(bold=True, size=16, color="1F3864")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


# --- Helpers -------------------------------------------------------------- #

def _write_header_row(
    sheet: Worksheet,
    columns: Iterable[ColumnSpec],
    *,
    start_row: int = 1,
) -> None:
    for idx, column in enumerate(columns, start=1):
        header_cell = sheet.cell(row=start_row, column=idx, value=column.header)
        header_cell.fill = _HEADER_FILL
        header_cell.font = _REQUIRED_FONT if column.required else _HEADER_FONT
        header_cell.alignment = _CENTER
        header_cell.border = _BORDER

        help_cell = sheet.cell(row=start_row + 1, column=idx, value=column.help_text)
        help_cell.fill = _SUBHEADER_FILL
        help_cell.font = _HELP_FONT
        help_cell.alignment = _WRAP
        help_cell.border = _BORDER

        sheet.column_dimensions[get_column_letter(idx)].width = column.width

    sheet.row_dimensions[start_row].height = 26
    sheet.row_dimensions[start_row + 1].height = 42
    sheet.freeze_panes = sheet.cell(row=start_row + 2, column=1)


def _write_catalog_sheet(sheet: Worksheet, bundle: CatalogBundle) -> dict[str, str]:
    """Vuelca los catalogos en columnas. Devuelve `nombre -> rango de codigos`."""
    sheet.cell(row=1, column=1, value="Catalogos controlados").font = _TITLE_FONT
    sheet.cell(row=2, column=1, value=(
        "Esta hoja contiene los valores aceptados por los dropdowns. "
        "Editar solo con autorizacion del equipo de arquitectura."
    )).font = _HELP_FONT

    ranges: dict[str, str] = {}
    catalogs = bundle.all()
    col_index = 1
    for name, catalog in catalogs.items():
        header = sheet.cell(row=4, column=col_index, value=catalog.name)
        header.fill = _HEADER_FILL
        header.font = _HEADER_FONT
        header.alignment = _CENTER
        for offset, entry in enumerate(catalog.entries, start=5):
            code_cell = sheet.cell(row=offset, column=col_index, value=entry.code)
            label_cell = sheet.cell(row=offset, column=col_index + 1, value=entry.label)
            code_cell.border = _BORDER
            label_cell.border = _BORDER
        sheet.column_dimensions[get_column_letter(col_index)].width = 22
        sheet.column_dimensions[get_column_letter(col_index + 1)].width = 32

        first_row = 5
        last_row = max(first_row, 4 + len(catalog.entries))
        code_letter = get_column_letter(col_index)
        ranges[name] = f"={sheet.title}!${code_letter}${first_row}:${code_letter}${last_row}"

        col_index += 3  # deja una columna en blanco entre catalogos

    sheet.freeze_panes = "A5"
    return ranges


def _add_validations(
    sheet: Worksheet,
    columns: Iterable[ColumnSpec],
    ranges: dict[str, str],
    *,
    data_start_row: int,
    data_end_row: int,
) -> None:
    for idx, column in enumerate(columns, start=1):
        col_letter = get_column_letter(idx)
        target_range = f"{col_letter}{data_start_row}:{col_letter}{data_end_row}"
        if column.kind == CellKind.DROPDOWN and column.catalog:
            formula = ranges.get(column.catalog)
            if not formula:
                continue
            dv = DataValidation(type="list", formula1=formula, allow_blank=not column.required)
            dv.error = f"Valor invalido para '{column.header}'. Use el dropdown."
            dv.errorTitle = "Valor no permitido"
            dv.prompt = column.help_text
            dv.promptTitle = column.header
            sheet.add_data_validation(dv)
            dv.add(target_range)
        elif column.kind == CellKind.DURATION:
            dv = DataValidation(
                type="custom",
                formula1=(
                    f'=OR(ISBLANK({col_letter}{data_start_row}),'
                    f'AND(LEN({col_letter}{data_start_row})>=2,'
                    f'OR(RIGHT({col_letter}{data_start_row},1)="h",'
                    f'RIGHT({col_letter}{data_start_row},1)="m",'
                    f'RIGHT({col_letter}{data_start_row},1)="s",'
                    f'RIGHT({col_letter}{data_start_row},1)="d")))'
                ),
                allow_blank=True,
            )
            dv.error = "Formato esperado: numero + unidad (ej. 2h, 30m, 1d)."
            dv.errorTitle = "SLA invalido"
            sheet.add_data_validation(dv)
            dv.add(target_range)
        elif column.kind == CellKind.ID and column.required:
            # Validacion de no-vacio (requiere texto >=3 caracteres).
            dv = DataValidation(
                type="textLength",
                operator="greaterThanOrEqual",
                formula1="3",
                allow_blank=False,
            )
            dv.error = "Los IDs deben tener al menos 3 caracteres (ej. PROC-001)."
            dv.errorTitle = "ID requerido"
            sheet.add_data_validation(dv)
            dv.add(target_range)


def _write_instructions(sheet: Worksheet, schema: ExcelSchema) -> None:
    sheet.cell(row=1, column=1, value="Plantilla BPMN Empresarial").font = _TITLE_FONT
    intro = (
        "Esta plantilla es la entrada oficial de la plataforma BPMN.\n"
        "Reglas obligatorias:\n"
        "  1. Completar primero la hoja 'Proceso'.\n"
        "  2. Una fila = una actividad/evento/compuerta en la hoja 'Actividades'.\n"
        "  3. Los IDs deben ser unicos (formato sugerido: PROC-001, ACT-001).\n"
        "  4. 'Tipo BPMN', 'Sistema', 'Responsable', 'Riesgo' y 'Control' usan dropdowns;\n"
        "     no escriba valores libres en esas columnas.\n"
        "  5. Nombre de actividad: 'Verbo + Objeto [+ Contexto]'. Ej. 'Validar identidad cliente'.\n"
        "  6. Regla de comandos: Actividad -> Sistema -> Comando.\n"
        "     El comando se invoca desde el Sistema indicado en la misma fila.\n"
        "  7. SLA en formato 'numero + unidad'. Ej. 2h, 30m, 1d.\n"
        "  8. La hoja 'Catalogos' es de solo lectura. Editar solo con autorizacion.\n"
    )
    cell = sheet.cell(row=3, column=1, value=intro)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = Font(size=11)
    sheet.column_dimensions["A"].width = 110
    sheet.row_dimensions[3].height = 220

    sheet.cell(row=6, column=1, value="Hojas de trabajo").font = Font(bold=True, size=12)
    listing = [
        f"  * {schema.process_sheet}: metadata del proceso (1 fila).",
        f"  * {schema.activities_sheet}: actividades/eventos/compuertas del proceso.",
        f"  * {schema.catalogs_sheet}: catalogos controlados (no editar).",
    ]
    for offset, line in enumerate(listing, start=7):
        sheet.cell(row=offset, column=1, value=line).font = Font(size=11)


# --- API publica ---------------------------------------------------------- #

def build_template(
    output_path: Path,
    *,
    schema: ExcelSchema | None = None,
    catalogs: CatalogBundle | None = None,
    activity_rows: int = 50,
) -> Path:
    """Construye la plantilla Excel oficial y la escribe en disco.

    Args:
        output_path: ruta destino del .xlsx.
        schema: schema de la plantilla (default: `build_default_schema()`).
        catalogs: catalogos cargados (default: cargados desde `catalogs/`).
        activity_rows: numero de filas con validaciones precargadas.
    """
    schema = schema or build_default_schema()
    catalogs = catalogs or load_catalogs()

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    default_sheet = wb.active
    default_sheet.title = schema.instructions_sheet
    _write_instructions(default_sheet, schema)

    process_sheet = wb.create_sheet(schema.process_sheet)
    _write_header_row(process_sheet, schema.process_columns)

    activities_sheet = wb.create_sheet(schema.activities_sheet)
    _write_header_row(activities_sheet, schema.activity_columns)

    catalogs_sheet = wb.create_sheet(schema.catalogs_sheet)
    catalog_ranges = _write_catalog_sheet(catalogs_sheet, catalogs)

    process_data_start = 3
    _add_validations(
        process_sheet,
        schema.process_columns,
        catalog_ranges,
        data_start_row=process_data_start,
        data_end_row=process_data_start + 2,
    )

    activity_data_start = 3
    _add_validations(
        activities_sheet,
        schema.activity_columns,
        catalog_ranges,
        data_start_row=activity_data_start,
        data_end_row=activity_data_start + activity_rows - 1,
    )

    wb.save(output_path)
    log.info("Plantilla Excel generada en {}", output_path)
    return output_path


# --- CLI ------------------------------------------------------------------ #

def cli(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="bpmn-template",
        description="Genera la plantilla Excel oficial de la plataforma BPMN.",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=settings.exports_dir / "plantilla_bpmn.xlsx",
        help="Ruta destino del archivo Excel.",
    )
    parser.add_argument(
        "--rows",
        type=int,
        default=50,
        help="Numero de filas con validaciones precargadas (default: 50).",
    )
    args = parser.parse_args(argv)
    settings.ensure_dirs()
    build_template(args.output, activity_rows=args.rows)
    print(f"OK: plantilla escrita en {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(cli())
