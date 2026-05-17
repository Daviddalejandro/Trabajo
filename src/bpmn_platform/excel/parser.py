"""Parser de la plantilla Excel oficial (Fase 4).

Convierte un archivo .xlsx generado con `build_template` en un
`EnterpriseModel` validado, recolectando issues (errores, warnings, infos)
en un `ParseResult`. Aplica:

  * Validacion de IDs unicos.
  * Resolucion de catalogos (codigo o etiqueta) -> entidades del modelo.
  * Regla 'Actividad -> Sistema -> Comando' (los comandos pertenecen
    a la aplicacion; un comando sin sistema valido es error).
  * Reglas de nomenclatura via `core.governance.validate_activity_name`.
  * Construccion de activos de informacion, KPIs y SLAs a partir de las
    columnas correspondientes.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..core.catalogs import CatalogBundle, load_catalogs
from ..core.catalogs.loader import Catalog, CatalogEntry
from ..core.governance import IssueSeverity, validate_activity_name
from ..core.metamodel import (
    KPI,
    SLA,
    Activity,
    ActivityApplicationUse,
    ActivityAssetFlow,
    Application,
    AssetFlowDirection,
    BpmnType,
    Command,
    Control,
    EnterpriseModel,
    Event,
    Gateway,
    GatewayKind,
    InformationAsset,
    Process,
    Risk,
    Role,
    Severity,
)
from ..logging_config import get_logger
from .schema import ExcelSchema, build_default_schema

log = get_logger(__name__)

_EVENT_TYPES = {
    BpmnType.START_EVENT,
    BpmnType.END_EVENT,
    BpmnType.INTERMEDIATE_CATCH_EVENT,
    BpmnType.MESSAGE_EVENT,
}
_GATEWAY_TYPES = {BpmnType.EXCLUSIVE_GATEWAY, BpmnType.PARALLEL_GATEWAY}


@dataclass
class ParseIssue:
    severity: IssueSeverity
    code: str
    message: str
    sheet: str | None = None
    row: int | None = None
    column: str | None = None

    @property
    def location(self) -> str:
        parts: list[str] = []
        if self.sheet:
            parts.append(self.sheet)
        if self.row is not None:
            parts.append(f"fila {self.row}")
        if self.column:
            parts.append(self.column)
        return " / ".join(parts) if parts else "-"


@dataclass
class ParseResult:
    model: EnterpriseModel
    issues: list[ParseIssue] = field(default_factory=list)

    @property
    def errors(self) -> list[ParseIssue]:
        return [i for i in self.issues if i.severity == IssueSeverity.ERROR]

    @property
    def warnings(self) -> list[ParseIssue]:
        return [i for i in self.issues if i.severity == IssueSeverity.WARNING]

    @property
    def infos(self) -> list[ParseIssue]:
        return [i for i in self.issues if i.severity == IssueSeverity.INFO]

    @property
    def ok(self) -> bool:
        return not self.errors


def _slugify(value: str) -> str:
    cleaned = re.sub(r"[^\w]+", "_", value or "", flags=re.UNICODE)
    return cleaned.strip("_").upper()


def _resolve_entry(catalog: Catalog, value: str | None) -> CatalogEntry | None:
    if not value:
        return None
    needle = value.strip().lower()
    for entry in catalog.entries:
        if entry.code.lower() == needle or entry.label.lower() == needle:
            return entry
    return None


class ExcelParser:
    """Parser de la plantilla oficial."""

    DATA_START_ROW = 3

    def __init__(
        self,
        schema: ExcelSchema | None = None,
        catalogs: CatalogBundle | None = None,
    ) -> None:
        self.schema = schema or build_default_schema()
        self.catalogs = catalogs or load_catalogs()
        self._issues: list[ParseIssue] = []
        self._asset_ids: dict[str, str] = {}
        self._command_ids: dict[tuple[str, str], str] = {}
        self._kpi_ids: dict[str, str] = {}

    # ------------------------------------------------------------------ #
    # API publica
    # ------------------------------------------------------------------ #

    def parse(self, path: Path | str) -> ParseResult:
        path = Path(path)
        log.info("Parseando Excel: {}", path)
        self._issues = []
        self._asset_ids = {}
        self._command_ids = {}
        self._kpi_ids = {}

        wb = load_workbook(path, data_only=True)

        for sheet_name in (self.schema.process_sheet, self.schema.activities_sheet):
            if sheet_name not in wb.sheetnames:
                self._add(
                    IssueSeverity.ERROR,
                    "SHEET-MISSING",
                    f"Falta la hoja obligatoria '{sheet_name}'.",
                )
                return ParseResult(model=EnterpriseModel(), issues=list(self._issues))

        seed = self._seed_from_catalogs()
        processes = self._parse_processes(wb[self.schema.process_sheet])
        activity_payload = self._parse_activities(wb[self.schema.activities_sheet])

        process_ids = {p.id for p in processes}
        for activity in activity_payload["activities"]:
            if activity.process_id not in process_ids:
                self._add(
                    IssueSeverity.ERROR,
                    "ACT-PROC-UNKNOWN",
                    f"La actividad '{activity.id}' referencia un proceso inexistente "
                    f"('{activity.process_id}').",
                )

        try:
            model = EnterpriseModel(processes=processes, **seed, **activity_payload)
        except Exception as exc:  # noqa: BLE001
            log.exception("Fallo construyendo EnterpriseModel.")
            self._add(IssueSeverity.ERROR, "MODEL-INVALID", str(exc))
            model = EnterpriseModel()

        return ParseResult(model=model, issues=list(self._issues))

    # ------------------------------------------------------------------ #
    # Semilla desde catalogos
    # ------------------------------------------------------------------ #

    def _seed_from_catalogs(self) -> dict[str, list]:
        roles = [
            Role(
                id=e.code,
                name=e.label,
                description=e.description,
                is_human=e.code != "SISTEMA",
            )
            for e in self.catalogs.roles.entries
        ]
        applications = [
            Application(id=e.code, name=e.label, description=e.description)
            for e in self.catalogs.systems.entries
        ]
        risks: list[Risk] = []
        for entry in self.catalogs.risks.entries:
            severity = Severity(getattr(entry, "severity", "medium"))
            risks.append(
                Risk(
                    id=entry.code,
                    name=entry.label,
                    description=entry.description,
                    severity=severity,
                )
            )
        controls: list[Control] = []
        for entry in self.catalogs.controls.entries:
            mitigates = getattr(entry, "mitigates", None) or []
            controls.append(
                Control(
                    id=entry.code,
                    name=entry.label,
                    description=entry.description,
                    mitigates_risk_ids=list(mitigates),
                    type=getattr(entry, "type", None),
                )
            )
        return {
            "roles": roles,
            "applications": applications,
            "risks": risks,
            "controls": controls,
        }

    # ------------------------------------------------------------------ #
    # Procesos
    # ------------------------------------------------------------------ #

    def _parse_processes(self, sheet: Worksheet) -> list[Process]:
        processes: list[Process] = []
        seen_ids: set[str] = set()
        cols = self.schema.process_columns

        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=self.DATA_START_ROW, max_row=sheet.max_row),
            start=self.DATA_START_ROW,
        ):
            values = [cell.value for cell in row]
            if not any(values):
                continue
            row_dict = self._row_to_dict(values, cols)
            process_id = self._clean_str(row_dict.get("process_id"))
            name = self._clean_str(row_dict.get("process_name"))

            if not process_id:
                self._add(
                    IssueSeverity.ERROR,
                    "PROC-ID-MISSING",
                    "Falta el ID del proceso.",
                    sheet.title,
                    row_idx,
                    "ID Proceso",
                )
                continue
            if process_id in seen_ids:
                self._add(
                    IssueSeverity.ERROR,
                    "PROC-ID-DUPLICATE",
                    f"ID de proceso duplicado: {process_id}.",
                    sheet.title,
                    row_idx,
                    "ID Proceso",
                )
                continue
            seen_ids.add(process_id)

            if not name:
                self._add(
                    IssueSeverity.ERROR,
                    "PROC-NAME-MISSING",
                    "Falta el nombre del proceso.",
                    sheet.title,
                    row_idx,
                    "Nombre del Proceso",
                )
                continue

            owner_value = self._clean_str(row_dict.get("owner_role"))
            owner_id: str | None = None
            if owner_value:
                entry = _resolve_entry(self.catalogs.roles, owner_value)
                if entry:
                    owner_id = entry.code
                else:
                    self._add(
                        IssueSeverity.WARNING,
                        "PROC-ROLE-UNKNOWN",
                        f"Rol '{owner_value}' no esta en el catalogo.",
                        sheet.title,
                        row_idx,
                        "Owner (Rol)",
                    )

            processes.append(
                Process(
                    id=process_id,
                    name=name,
                    domain=self._clean_str(row_dict.get("domain")),
                    owner_role_id=owner_id,
                    version=self._clean_str(row_dict.get("version")) or "1.0.0",
                    description=self._clean_str(row_dict.get("description")),
                )
            )

        if not processes:
            self._add(
                IssueSeverity.ERROR,
                "PROC-EMPTY",
                f"La hoja '{sheet.title}' no contiene procesos.",
            )
        return processes

    # ------------------------------------------------------------------ #
    # Actividades / eventos / compuertas
    # ------------------------------------------------------------------ #

    def _parse_activities(self, sheet: Worksheet) -> dict[str, list]:
        activities: list[Activity] = []
        events: list[Event] = []
        gateways: list[Gateway] = []
        commands: list[Command] = []
        info_assets: list[InformationAsset] = []
        kpis: list[KPI] = []
        slas: list[SLA] = []
        app_uses: list[ActivityApplicationUse] = []
        asset_flows: list[ActivityAssetFlow] = []

        seen_ids: set[str] = set()
        cols = self.schema.activity_columns

        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=self.DATA_START_ROW, max_row=sheet.max_row),
            start=self.DATA_START_ROW,
        ):
            values = [cell.value for cell in row]
            if not any(values):
                continue
            row_dict = self._row_to_dict(values, cols)

            act_id = self._clean_str(row_dict.get("activity_id"))
            process_id = self._clean_str(row_dict.get("process_id"))
            bpmn_value = self._clean_str(row_dict.get("bpmn_type"))
            name = self._clean_str(row_dict.get("activity_name"))

            if not act_id:
                self._add(
                    IssueSeverity.ERROR,
                    "ACT-ID-MISSING",
                    "Falta el ID de la actividad/evento.",
                    sheet.title,
                    row_idx,
                    "ID",
                )
                continue
            if act_id in seen_ids:
                self._add(
                    IssueSeverity.ERROR,
                    "ACT-ID-DUPLICATE",
                    f"ID duplicado: {act_id}.",
                    sheet.title,
                    row_idx,
                    "ID",
                )
                continue
            seen_ids.add(act_id)

            if not process_id:
                self._add(
                    IssueSeverity.ERROR,
                    "ACT-PROC-MISSING",
                    "Falta el ID del proceso al que pertenece la actividad.",
                    sheet.title,
                    row_idx,
                    "ID Proceso",
                )
                continue

            bpmn_type = self._resolve_bpmn_type(bpmn_value, row_idx, sheet.title)
            if bpmn_type is None:
                continue

            if not name:
                self._add(
                    IssueSeverity.ERROR,
                    "ACT-NAME-MISSING",
                    "Falta el nombre de la actividad/evento.",
                    sheet.title,
                    row_idx,
                    "Actividad",
                )
                continue

            # Naming rule solo aplica a tasks (no eventos/gateways).
            if bpmn_type not in _EVENT_TYPES and bpmn_type not in _GATEWAY_TYPES:
                for issue in validate_activity_name(name, location=f"{sheet.title}!{act_id}"):
                    self._add(issue.severity, issue.code, issue.message, sheet.title, row_idx, "Actividad")

            if bpmn_type in _EVENT_TYPES:
                events.append(
                    Event(
                        id=act_id,
                        name=name,
                        bpmn_type=bpmn_type,
                        is_start=bpmn_type == BpmnType.START_EVENT,
                        is_end=bpmn_type == BpmnType.END_EVENT,
                    )
                )
                continue

            if bpmn_type in _GATEWAY_TYPES:
                kind = (
                    GatewayKind.EXCLUSIVE
                    if bpmn_type == BpmnType.EXCLUSIVE_GATEWAY
                    else GatewayKind.PARALLEL
                )
                gateways.append(Gateway(id=act_id, name=name, kind=kind))
                continue

            # ------------------- UserTask / ServiceTask -------------------

            role_id = self._resolve_catalog_value(
                row_dict.get("role"), self.catalogs.roles,
                row_idx, sheet.title, "Responsable", warning_code="ACT-ROLE-UNKNOWN",
            )

            app_id = self._resolve_catalog_value(
                row_dict.get("system"), self.catalogs.systems,
                row_idx, sheet.title, "Sistema", warning_code="ACT-SYSTEM-UNKNOWN",
            )

            command_value = self._clean_str(row_dict.get("command"))
            command_id: str | None = None
            if command_value:
                if not app_id:
                    self._add(
                        IssueSeverity.ERROR,
                        "CMD-ORPHAN",
                        f"Se especifico el comando '{command_value}' sin un Sistema valido. "
                        "Regla obligatoria: Actividad -> Sistema -> Comando.",
                        sheet.title,
                        row_idx,
                        "Comando",
                    )
                else:
                    key = (app_id, command_value)
                    cmd_id = self._command_ids.get(key)
                    if not cmd_id:
                        cmd_id = self._unique_id(
                            f"CMD-{app_id}-{_slugify(command_value)[:30]}",
                            {c.id for c in commands},
                        )
                        commands.append(
                            Command(
                                id=cmd_id,
                                name=command_value,
                                application_id=app_id,
                                invocation=command_value,
                            )
                        )
                        self._command_ids[key] = cmd_id
                    command_id = cmd_id

            risk_id = self._resolve_catalog_value(
                row_dict.get("risk"), self.catalogs.risks,
                row_idx, sheet.title, "Riesgo", warning_code="ACT-RISK-UNKNOWN",
            )

            control_value = self._clean_str(row_dict.get("control"))
            control_id: str | None = None
            if control_value:
                entry = _resolve_entry(self.catalogs.controls, control_value)
                if entry:
                    control_id = entry.code
                    mitigates = getattr(entry, "mitigates", None) or []
                    if risk_id and risk_id not in mitigates:
                        self._add(
                            IssueSeverity.INFO,
                            "CTRL-RISK-MISMATCH",
                            f"El control '{control_id}' no declara mitigacion del riesgo '{risk_id}'.",
                            sheet.title,
                            row_idx,
                            "Control",
                        )
                else:
                    self._add(
                        IssueSeverity.WARNING,
                        "ACT-CONTROL-UNKNOWN",
                        f"Control '{control_value}' no esta en catalogo.",
                        sheet.title,
                        row_idx,
                        "Control",
                    )

            sla_value = self._clean_str(row_dict.get("sla"))
            sla_id: str | None = None
            if sla_value:
                sla_id = f"SLA-{act_id}"
                try:
                    slas.append(SLA(id=sla_id, name=f"SLA {name}", duration=sla_value))
                except Exception as exc:  # noqa: BLE001
                    self._add(
                        IssueSeverity.ERROR,
                        "SLA-INVALID",
                        f"SLA invalido '{sla_value}': {exc}",
                        sheet.title,
                        row_idx,
                        "SLA",
                    )
                    sla_id = None

            kpi_value = self._clean_str(row_dict.get("kpi"))
            kpi_id: str | None = None
            if kpi_value:
                kpi_id = self._kpi_ids.get(kpi_value)
                if not kpi_id:
                    kpi_id = self._unique_id(
                        f"KPI-{_slugify(kpi_value)[:30]}", {k.id for k in kpis}
                    )
                    self._kpi_ids[kpi_value] = kpi_id
                    kpis.append(KPI(id=kpi_id, name=kpi_value))

            for column_key, direction in (
                ("input", AssetFlowDirection.CONSUMES),
                ("output", AssetFlowDirection.PRODUCES),
                ("information_asset", AssetFlowDirection.CONSUMES),
            ):
                value = self._clean_str(row_dict.get(column_key))
                if not value:
                    continue
                asset_id = self._asset_ids.get(value)
                if not asset_id:
                    asset_id = self._unique_id(
                        f"INFO-{_slugify(value)[:30]}", set(self._asset_ids.values())
                    )
                    self._asset_ids[value] = asset_id
                    info_assets.append(InformationAsset(id=asset_id, name=value))
                asset_flows.append(
                    ActivityAssetFlow(
                        activity_id=act_id, asset_id=asset_id, direction=direction
                    )
                )

            observations = self._clean_str(row_dict.get("observations"))

            activities.append(
                Activity(
                    id=act_id,
                    name=name,
                    bpmn_type=bpmn_type,
                    process_id=process_id,
                    role_id=role_id,
                    application_ids=[app_id] if app_id else [],
                    risk_ids=[risk_id] if risk_id else [],
                    control_ids=[control_id] if control_id else [],
                    kpi_ids=[kpi_id] if kpi_id else [],
                    sla_id=sla_id,
                    observations=observations,
                )
            )

            if app_id:
                app_uses.append(
                    ActivityApplicationUse(
                        activity_id=act_id,
                        application_id=app_id,
                        command_ids=[command_id] if command_id else [],
                    )
                )

        if not activities and not events and not gateways:
            self._add(
                IssueSeverity.ERROR,
                "ACT-EMPTY",
                f"La hoja '{sheet.title}' no contiene actividades.",
            )

        return {
            "activities": activities,
            "events": events,
            "gateways": gateways,
            "commands": commands,
            "information_assets": info_assets,
            "kpis": kpis,
            "slas": slas,
            "activity_application_uses": app_uses,
            "activity_asset_flows": asset_flows,
        }

    # ------------------------------------------------------------------ #
    # Helpers
    # ------------------------------------------------------------------ #

    def _resolve_bpmn_type(
        self, value: str | None, row_idx: int, sheet_name: str
    ) -> BpmnType | None:
        if not value:
            self._add(
                IssueSeverity.ERROR,
                "BPMN-TYPE-MISSING",
                "Falta el tipo BPMN.",
                sheet_name,
                row_idx,
                "Tipo BPMN",
            )
            return None
        text = value.strip()
        for bpmn in BpmnType:
            if bpmn.value.lower() == text.lower():
                return bpmn
        for label, bpmn in BpmnType.business_labels().items():
            if label.lower() == text.lower():
                return bpmn
        self._add(
            IssueSeverity.ERROR,
            "BPMN-TYPE-INVALID",
            f"Tipo BPMN no reconocido: '{text}'.",
            sheet_name,
            row_idx,
            "Tipo BPMN",
        )
        return None

    def _resolve_catalog_value(
        self,
        raw_value: Any,
        catalog: Catalog,
        row_idx: int,
        sheet_name: str,
        column_label: str,
        *,
        warning_code: str,
    ) -> str | None:
        value = self._clean_str(raw_value)
        if not value:
            return None
        entry = _resolve_entry(catalog, value)
        if entry:
            return entry.code
        self._add(
            IssueSeverity.WARNING,
            warning_code,
            f"Valor '{value}' no esta en el catalogo '{catalog.name}'.",
            sheet_name,
            row_idx,
            column_label,
        )
        return None

    @staticmethod
    def _row_to_dict(values: list[Any], cols) -> dict[str, Any]:
        return {
            col.key: values[idx] if idx < len(values) else None
            for idx, col in enumerate(cols)
        }

    @staticmethod
    def _clean_str(value: Any) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _unique_id(base: str, used: set[str]) -> str:
        if base not in used:
            return base
        counter = 2
        while f"{base}-{counter}" in used:
            counter += 1
        return f"{base}-{counter}"

    def _add(
        self,
        severity: IssueSeverity,
        code: str,
        message: str,
        sheet: str | None = None,
        row: int | None = None,
        column: str | None = None,
    ) -> None:
        self._issues.append(
            ParseIssue(
                severity=severity,
                code=code,
                message=message,
                sheet=sheet,
                row=row,
                column=column,
            )
        )
