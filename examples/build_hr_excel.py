"""Excel diligenciado de muestra: 'Vinculacion de nuevo talento humano'.

Uso:
    python -m examples.build_hr_excel [-o RUTA.xlsx]
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from bpmn_platform.config import settings
from bpmn_platform.excel.schema import build_default_schema
from bpmn_platform.excel.template_builder import build_template


def _process_row() -> tuple:
    # process_id, process_name, domain, owner_role, version, description
    return (
        "PROC-VINCULACION-TH",
        "Vinculacion de nuevo talento humano",
        "Talento Humano",
        "COORD_TALENTO_HUMANO",
        "1.0.0",
        (
            "Proceso integral para la vinculacion de un nuevo colaborador, "
            "desde la recepcion de la vacante hasta la induccion corporativa."
        ),
    )


def _activity_rows() -> list[tuple]:
    # Orden de columnas (segun build_default_schema):
    #  1 activity_id      2 process_id
    #  3 bpmn_type        4 activity_name
    #  5 role             6 system          7 command         8 api
    #  9 input            10 output         11 information_asset
    # 12 risk             13 control        14 sla            15 kpi
    # 16 observations     17 predecessors
    return [
        # ---- Inicio
        (
            "EVT-START",
            "PROC-VINCULACION-TH",
            "Inicio",
            "Requisicion de vacante recibida",
            None, None, None, None, None, None, None,
            None, None, None, None,
            "Inicia con la requisicion firmada por el lider del area.",
        ),
        # ---- Publicar vacante (automatico)
        (
            "ACT-PUBLICAR",
            "PROC-VINCULACION-TH",
            "ServiceTask",
            "Publicar vacante en portales",
            "SISTEMA",
            "PORTAL_EMPLEO",
            "publish_vacancy --requisition $REQ_ID",
            "/v1/jobs/publish",
            "Requisicion de vacante",
            "Vacante publicada",
            "Requisicion de vacante",
            None,
            "C-LOG_AUDITORIA",
            "4h",
            "Tiempo medio de publicacion",
            "Publicacion automatica en LinkedIn y elempleo.",
        ),
        # ---- Filtrar HVs
        (
            "ACT-FILTRAR",
            "PROC-VINCULACION-TH",
            "Actividad Humana",
            "Filtrar hojas de vida",
            "ANALISTA_TALENTO_HUMANO",
            "HRM",
            None, None,
            "Hojas de vida recibidas",
            "Lista corta de candidatos",
            "Hojas de vida recibidas",
            "R-ERROR_OPERATIVO",
            "C-DOBLE_VALIDACION",
            "2d",
            "Tasa de candidatos filtrados",
            "Verificar competencias minimas vs requisicion.",
        ),
        # ---- Entrevista tecnica
        (
            "ACT-ENTREVISTA",
            "PROC-VINCULACION-TH",
            "Actividad Humana",
            "Realizar entrevista tecnica",
            "LIDER_AREA",
            "HRM",
            None, None,
            "Lista corta de candidatos",
            "Resultado entrevista",
            "Hojas de vida recibidas",
            "R-ERROR_OPERATIVO",
            "C-DOBLE_VALIDACION",
            "1d",
            "Calificacion promedio de entrevistas",
            "Entrevista conjunta TH + Lider del area.",
        ),
        # ---- Aplicar pruebas psicotecnicas
        (
            "ACT-PSICOTECNICAS",
            "PROC-VINCULACION-TH",
            "ServiceTask",
            "Aplicar pruebas psicotecnicas",
            "SISTEMA",
            "TEST_PSICOTECNICOS",
            "python run_assessment.py --candidate $CAND_ID",
            "/v1/assessments/launch",
            "Lista corta de candidatos",
            "Resultados psicotecnicos",
            "Resultados psicotecnicos",
            "R-PERDIDA_INFO",
            "C-LOG_AUDITORIA",
            "1d",
            "Tasa de pruebas completadas",
            "Plataforma externa con SSO; resultados consumidos por API.",
        ),
        # ---- Gateway: candidato apto?
        (
            "GW-APTO",
            "PROC-VINCULACION-TH",
            "Decision",
            "Candidato apto?",
            None, None, None, None, None, None, None,
            None, None, None, None,
            "Decision sobre continuar o no con el candidato.",
        ),
        # ---- Verificar referencias y antecedentes (rama SI)
        (
            "ACT-VERIFICAR",
            "PROC-VINCULACION-TH",
            "Actividad Humana",
            "Verificar referencias y antecedentes",
            "ANALISTA_TALENTO_HUMANO",
            "HRM",
            None, None,
            "Resultados psicotecnicos",
            "Verificacion completa",
            "Documento identidad",
            "R-SUPLANTACION",
            "C-BIOMETRICO",
            "2d",
            "Tasa de verificaciones exitosas",
            "Validacion biometrica + chequeo de listas restrictivas.",
            "GW-APTO",
        ),
        # ---- Generar oferta laboral
        (
            "ACT-OFERTA",
            "PROC-VINCULACION-TH",
            "Actividad Humana",
            "Generar y enviar oferta laboral",
            "COORD_TALENTO_HUMANO",
            "HRM",
            None, None,
            "Verificacion completa",
            "Oferta laboral enviada",
            "Oferta laboral",
            "R-FRAUDE",
            "C-DOBLE_VALIDACION",
            "1d",
            "Tasa de ofertas aceptadas",
            "Aprobacion previa del Lider del Area + Coordinador TH.",
            "ACT-VERIFICAR",
        ),
        # ---- Firmar contrato laboral
        (
            "ACT-CONTRATO",
            "PROC-VINCULACION-TH",
            "Actividad Humana",
            "Firmar contrato laboral",
            "COORD_TALENTO_HUMANO",
            "HRM",
            None, None,
            "Oferta laboral aceptada",
            "Contrato laboral firmado",
            "Contrato laboral",
            "R-INCUMPLIMIENTO",
            "C-CHECKLIST_REG",
            "2d",
            "Cumplimiento legal de contratacion",
            "Firma electronica + checklist de documentos obligatorios.",
            "ACT-OFERTA",
        ),
        # ---- Crear usuarios en sistemas
        (
            "ACT-CREAR-USR",
            "PROC-VINCULACION-TH",
            "ServiceTask",
            "Crear usuarios en sistemas corporativos",
            "SISTEMA",
            "ACTIVE_DIRECTORY",
            "python provision_user.py --employee $EMP_ID",
            "/v1/iam/provision",
            "Contrato laboral firmado",
            "Usuarios creados",
            "Credenciales corporativas",
            "R-PERDIDA_INFO",
            "C-LOG_AUDITORIA",
            "4h",
            "Tiempo medio de aprovisionamiento",
            "Provisioning automatico via Airflow + AD.",
            "ACT-CONTRATO",
        ),
        # ---- Induccion corporativa
        (
            "ACT-INDUCCION",
            "PROC-VINCULACION-TH",
            "Actividad Humana",
            "Realizar induccion corporativa",
            "COORD_TALENTO_HUMANO",
            "HRM",
            None, None,
            "Usuarios creados",
            "Empleado inducido",
            "Plan de induccion",
            "R-DISPONIBILIDAD",
            "C-MONITOREO_SLA",
            "5d",
            "Cobertura de induccion",
            "Incluye SST, codigo de etica y conocimiento del producto.",
            "ACT-CREAR-USR",
        ),
        # ---- Fin OK
        (
            "EVT-END-OK",
            "PROC-VINCULACION-TH",
            "Fin",
            "Colaborador activo",
            None, None, None, None, None, None, None,
            None, None, None, None,
            "Empleado en nomina y operando.",
            "ACT-INDUCCION",
        ),
        # ---- Notificar no continuacion (rama NO)
        (
            "ACT-NO-CONTINUAR",
            "PROC-VINCULACION-TH",
            "ServiceTask",
            "Notificar candidato no seleccionado",
            "SISTEMA",
            "EMAIL",
            "send_email --template=hr_rejection",
            None,
            "Resultados psicotecnicos",
            "Carta rechazo",
            None,
            "R-PERDIDA_INFO",
            "C-LOG_AUDITORIA",
            "1d",
            "Tiempo medio de respuesta al candidato",
            "Email automatico personalizado.",
            "GW-APTO",
        ),
        # ---- Fin KO
        (
            "EVT-END-KO",
            "PROC-VINCULACION-TH",
            "Fin",
            "Candidato no seleccionado",
            None, None, None, None, None, None, None,
            None, None, None, None,
            "Cierre del flujo para el candidato.",
            "ACT-NO-CONTINUAR",
        ),
    ]


def build_sample(output_path: Path) -> Path:
    schema = build_default_schema()
    build_template(output_path, schema=schema, activity_rows=40)
    wb = load_workbook(output_path)
    process_sheet = wb[schema.process_sheet]
    for col, value in enumerate(_process_row(), start=1):
        process_sheet.cell(row=3, column=col, value=value)
    activities_sheet = wb[schema.activities_sheet]
    for r, row in enumerate(_activity_rows(), start=3):
        for col, value in enumerate(row, start=1):
            activities_sheet.cell(row=r, column=col, value=value)
    wb.save(output_path)
    return output_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="build-hr-excel")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=settings.exports_dir / "vinculacion_talento_humano.xlsx",
    )
    args = parser.parse_args(argv)
    settings.ensure_dirs()
    path = build_sample(args.output)
    print(f"OK: Excel HR generado en {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
