"""Genera un Excel de muestra ya diligenciado: 'Onboarding de clientes'.

Uso:
    python -m examples.build_sample_excel [-o RUTA.xlsx]
"""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from bpmn_platform.config import settings
from bpmn_platform.excel.schema import build_default_schema
from bpmn_platform.excel.template_builder import build_template


def _process_row() -> tuple:
    # Columnas hoja Proceso: process_id, process_name, domain, owner_role, version, description
    return (
        "PROC-ONBOARDING",
        "Onboarding de clientes",
        "Comercial / Riesgos",
        "USUARIO_NEGOCIO",
        "1.0.0",
        "Alta integral de un cliente nuevo, desde solicitud hasta activacion.",
    )


def _activity_rows() -> list[tuple]:
    # Orden de columnas de la hoja Actividades (segun build_default_schema):
    # 1  activity_id        2  process_id
    # 3  bpmn_type          4  activity_name
    # 5  role               6  system          7  command       8  api
    # 9  input              10 output          11 information_asset
    # 12 risk               13 control          14 sla          15 kpi
    # 16 observations       17 predecessors
    return [
        # Start event
        ("EVT-START", "PROC-ONBOARDING", "Inicio", "Solicitud recibida"),
        # UserTask: validar identidad
        (
            "ACT-VALIDAR",
            "PROC-ONBOARDING",
            "Actividad Humana",
            "Validar identidad cliente",
            "ANALISTA_RIESGO",
            "CRM",
            None,
            None,
            "Documento identidad",
            "Identidad verificada",
            "Documento identidad",
            "R-SUPLANTACION",
            "C-BIOMETRICO",
            "2h",
            "Tasa de validacion exitosa",
            "Incluye validacion biometrica.",
        ),
        # ServiceTask: consultar listas restrictivas (automatico)
        (
            "ACT-LISTAS",
            "PROC-ONBOARDING",
            "ServiceTask",
            "Consultar listas restrictivas",
            "SISTEMA",
            "CUSTOM_API",
            "python check_lists.py",
            "/v1/sanctions/check",
            None,
            "Resultado AML",
            None,
            "R-INCUMPLIMIENTO",
            "C-CHECKLIST_REG",
            "30m",
            "Cobertura AML",
        ),
        # Gateway: aprobado?
        ("GW-DECISION", "PROC-ONBOARDING", "Decision", "Cliente aprobable?"),
        # UserTask: aprobar (rama si)
        (
            "ACT-APROBAR",
            "PROC-ONBOARDING",
            "Actividad Humana",
            "Aprobar solicitud cliente",
            "APROBADOR_CREDITO",
            "CRM",
            None,
            None,
            "Resultado AML",
            "Aprobacion final",
            None,
            "R-FRAUDE",
            "C-DOBLE_VALIDACION",
            "4h",
            "SLA aprobacion",
            "Ruta si: cliente aprobado.",
            "GW-DECISION",
        ),
        # ServiceTask: generar reporte regulatorio (despues de aprobar)
        (
            "ACT-REPORTE",
            "PROC-ONBOARDING",
            "ServiceTask",
            "Generar reporte regulatorio",
            "SISTEMA",
            "AIRFLOW",
            "python report_regulatory.py",
            None,
            None,
            "Reporte regulatorio",
            None,
            "R-INCUMPLIMIENTO",
            "C-LOG_AUDITORIA",
            "1h",
            "Reportes generados",
            None,
            "ACT-APROBAR",
        ),
        # Notificar al cliente
        (
            "ACT-NOTIFICAR",
            "PROC-ONBOARDING",
            "ServiceTask",
            "Notificar resultado al cliente",
            "SISTEMA",
            "EMAIL",
            "send_email --template=approval",
            None,
            None,
            "Email enviado",
            None,
            None,
            None,
            "1h",
            None,
            None,
            "ACT-REPORTE",
        ),
        # UserTask: rechazar (rama no)
        (
            "ACT-RECHAZAR",
            "PROC-ONBOARDING",
            "Actividad Humana",
            "Rechazar solicitud cliente",
            "APROBADOR_CREDITO",
            "CRM",
            None,
            None,
            "Resultado AML",
            "Carta rechazo",
            None,
            None,
            "C-LOG_AUDITORIA",
            "2h",
            None,
            "Ruta no: cliente rechazado.",
            "GW-DECISION",
        ),
        # End event: aprobado
        (
            "EVT-END-OK",
            "PROC-ONBOARDING",
            "Fin",
            "Cliente activo",
            None, None, None, None, None, None, None, None, None, None, None, None,
            "ACT-NOTIFICAR",
        ),
        # End event: rechazado
        (
            "EVT-END-KO",
            "PROC-ONBOARDING",
            "Fin",
            "Cliente rechazado",
            None, None, None, None, None, None, None, None, None, None, None, None,
            "ACT-RECHAZAR",
        ),
    ]


def build_sample(output_path: Path) -> Path:
    schema = build_default_schema()
    build_template(output_path, schema=schema, activity_rows=30)

    wb = load_workbook(output_path)
    process_sheet = wb[schema.process_sheet]
    for col, value in enumerate(_process_row(), start=1):
        process_sheet.cell(row=3, column=col, value=value)

    activities_sheet = wb[schema.activities_sheet]
    for r, row in enumerate(_activity_rows(), start=3):
        for col, value in enumerate(row, start=1):
            activities_sheet.cell(row=r, column=col, value=value)
    wb.save(output_path)
    return output_path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="build-sample-excel")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=settings.exports_dir / "sample_onboarding.xlsx",
    )
    args = parser.parse_args(argv)
    settings.ensure_dirs()
    path = build_sample(args.output)
    print(f"OK: Excel de muestra generado en {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
